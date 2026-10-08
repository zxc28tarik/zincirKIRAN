#!/usr/bin/env python3
from __future__ import annotations

import concurrent.futures
import gzip
import hashlib
import html as html_lib
import json
import io
import re
import time
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pandas as pd
import requests
import yfinance as yf
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from zincir_kiran.corporate_action_economic_resolution import (
    EconomicResolutionEvidence,
    resolve_economic_action,
)

ROOT = Path(__file__).resolve().parents[1]
HIST = (
    ROOT
    / "data/live_shadow/high52_historical_ca_reconciliation_v1"
    / "historical_ca_risks.csv"
)
RECENT = (
    ROOT
    / "data/live_shadow/high52_input_ca_risk_v1"
    / "recent_ca_risks.csv"
)
UNIVERSE = (
    ROOT
    / "data/live_shadow/first_post_activation_live_snapshot_v1"
    / "universe.csv"
)
OUT = ROOT / "data/live_shadow/corporate_action_economic_resolution_v1"
KAP_TEMPLATE = "https://www.kap.org.tr/tr/Bildirim/{event_id}"
KAP_EXCEL_TEMPLATE = "https://www.kap.org.tr/tr/api/notification/export/excel/{event_id}"
ACTIVATION = datetime.fromisoformat("2026-10-07T22:10:15+00:00")
PRICE_START = "2025-07-01"

PUSH_RE = re.compile(
    r'self\.__next_f\.push\(\[1,"((?:\\.|[^"\\])*)"\]\)</script>'
)
BASIC_RE = re.compile(
    r'"disclosureBasic":(\{.*?\}),"disclosureDetail":',
    re.DOTALL,
)


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def write_gzip(path: Path, payload: bytes) -> str:
    zipped = gzip.compress(payload, compresslevel=9, mtime=0)
    path.write_bytes(zipped)
    return sha256_bytes(zipped)


def load_queue() -> pd.DataFrame:
    hist = pd.read_csv(HIST, dtype=str, keep_default_na=False)
    recent = pd.read_csv(RECENT, dtype=str, keep_default_na=False)
    hist["scope"] = "historical"
    recent["scope"] = "recent"
    rows = pd.concat([hist, recent], ignore_index=True)

    required = {"ticker", "event_id", "event_type", "publish_date", "subject", "summary"}
    missing = required - set(rows.columns)
    if missing:
        raise RuntimeError(f"risk queue missing columns: {sorted(missing)}")

    grouped = []
    for event_id, group in rows.groupby("event_id", sort=True):
        event_types = sorted(set(group["event_type"].astype(str)))
        if len(event_types) != 1:
            raise RuntimeError(
                f"event {event_id} has conflicting types: {event_types}"
            )
        grouped.append(
            {
                "event_id": str(event_id),
                "event_type": event_types[0],
                "publish_date": sorted(set(group["publish_date"].astype(str)))[0],
                "subject": sorted(set(group["subject"].astype(str)))[0],
                "summary": sorted(set(group["summary"].astype(str)))[0],
                "tickers": "|".join(sorted(set(group["ticker"].astype(str)))),
                "scopes": "|".join(sorted(set(group["scope"].astype(str)))),
            }
        )
    queue = pd.DataFrame(grouped).sort_values("event_id").reset_index(drop=True)
    if len(queue) != 400:
        raise RuntimeError(f"expected 400 unique events, found {len(queue)}")
    return queue


def fetch_kap(event_id: str, attempts: int = 2) -> tuple[bytes | None, str | None]:
    url = KAP_TEMPLATE.format(event_id=event_id)
    headers = {
        "User-Agent": "zincir-kiran-ca-resolution/1.0",
        "Accept-Language": "tr-TR,tr;q=0.9",
    }
    last = None
    for attempt in range(1, attempts + 1):
        try:
            response = requests.get(url, headers=headers, timeout=20)
            if response.status_code == 429:
                retry_after = response.headers.get("Retry-After")
                try:
                    delay = min(max(float(retry_after or 0), 2.0), 30.0)
                except ValueError:
                    delay = 5.0
                last = "HTTP_429"
                time.sleep(delay)
                continue
            if response.status_code == 200 and response.content:
                raw = response.content
                text = raw.decode("utf-8", errors="ignore")
                if (
                    f'\"disclosureIndex\":{event_id}' not in text
                    and f'/Bildirim/{event_id}' not in text
                ):
                    last = "EVENT_ID_NOT_PRESENT_IN_PAGE"
                else:
                    return raw, None
            else:
                last = f"HTTP_{response.status_code}"
        except Exception as exc:  # noqa: BLE001
            last = f"{type(exc).__name__}:{exc}"
        time.sleep(min(2**attempt, 10))
    return None, last


def decode_next_payload(raw: bytes) -> str:
    text = raw.decode("utf-8", errors="ignore")
    chunks = []
    for match in PUSH_RE.finditer(text):
        encoded = match.group(1)
        try:
            decoded = json.loads('"' + encoded + '"')
        except json.JSONDecodeError:
            continue
        chunks.append(decoded)
    return "\n".join(chunks)


def parse_basic(expanded: str) -> dict:
    match = BASIC_RE.search(expanded)
    if not match:
        return {}
    try:
        return json.loads(match.group(1))
    except json.JSONDecodeError:
        return {}


def _taxonomy_fields_from_html(document: str) -> list[dict]:
    soup = BeautifulSoup(document, "html.parser")
    fields: list[dict] = []
    for row in soup.find_all("tr"):
        name_node = row.select_one(".taxonomy-field-name")
        if name_node is None:
            continue
        field_name = name_node.get_text(" ", strip=True).rstrip("|").strip()
        if not field_name:
            continue

        values: list[str] = []
        selectors = (
            ".taxonomy-context-value.content-tr",
            ".taxonomy-context-value-summernote.content-tr",
        )
        for selector in selectors:
            for node in row.select(selector):
                value = html_lib.unescape(node.get_text(" ", strip=True)).strip()
                if value and value not in values:
                    values.append(value)

        titles: list[str] = []
        for node in row.select(".taxonomy-field-title .content-tr"):
            value = html_lib.unescape(node.get_text(" ", strip=True)).strip()
            if value and value not in titles:
                titles.append(value)

        fields.append(
            {
                "field_name": field_name,
                "titles_tr": titles,
                "values_tr": values,
            }
        )
    return fields


def parse_taxonomy_fields(raw: bytes, expanded: str) -> list[dict]:
    """Merge taxonomy rows from raw SSR HTML and decoded Next payload."""
    combined: dict[tuple[str, tuple[str, ...], tuple[str, ...]], dict] = {}
    documents = [
        raw.decode("utf-8", errors="ignore"),
        expanded,
    ]
    for document in documents:
        for field in _taxonomy_fields_from_html(document):
            key = (
                field["field_name"],
                tuple(field["titles_tr"]),
                tuple(field["values_tr"]),
            )
            combined[key] = field
    return sorted(
        combined.values(),
        key=lambda item: (
            item["field_name"],
            item["titles_tr"],
            item["values_tr"],
        ),
    )


def fetch_kap_excel(
    event_id: str,
    attempts: int = 3,
) -> tuple[bytes | None, str | None]:
    url = KAP_EXCEL_TEMPLATE.format(event_id=event_id)
    headers = {
        "User-Agent": "zincir-kiran-ca-resolution/1.0",
        "Accept-Language": "tr-TR,tr;q=0.9",
        "Accept": (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet,application/octet-stream,*/*"
        ),
    }
    last = None
    for attempt in range(1, attempts + 1):
        try:
            response = requests.get(url, headers=headers, timeout=30)
            if response.status_code == 429:
                retry_after = response.headers.get("Retry-After")
                try:
                    delay = min(max(float(retry_after or 0), 2.0), 30.0)
                except ValueError:
                    delay = 5.0
                last = "HTTP_429"
                time.sleep(delay)
                continue
            if response.status_code == 200 and response.content:
                raw = response.content
                if raw[:2] != b"PK":
                    last = "NOT_XLSX_ZIP"
                else:
                    return raw, None
            else:
                last = f"HTTP_{response.status_code}"
        except Exception as exc:  # noqa: BLE001
            last = f"{type(exc).__name__}:{exc}"
        time.sleep(min(2**attempt, 10))
    return None, last


def parse_excel_cells(raw: bytes) -> list[dict]:
    workbook = load_workbook(
        io.BytesIO(raw),
        read_only=True,
        data_only=True,
    )
    rows: list[dict] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    if isinstance(value, str):
                        normalized = value.strip()
                        if not normalized:
                            continue
                        rendered = normalized
                    elif isinstance(value, datetime):
                        rendered = value.isoformat()
                    else:
                        rendered = str(value)
                    rows.append(
                        {
                            "sheet": sheet.title,
                            "coordinate": cell.coordinate,
                            "row": int(cell.row),
                            "column": int(cell.column),
                            "value": rendered,
                            "value_type": type(value).__name__,
                        }
                    )
    finally:
        workbook.close()
    return rows


def capture_excel_details(
    queue: pd.DataFrame,
) -> tuple[pd.DataFrame, list[dict]]:
    excel_dir = OUT / "kap_excel"
    parsed_dir = OUT / "parsed_excel"
    excel_dir.mkdir(parents=True, exist_ok=True)
    parsed_dir.mkdir(parents=True, exist_ok=True)

    rows: list[dict] = []
    cell_records: list[dict] = []

    def one(event_id: str) -> dict:
        path = excel_dir / f"{event_id}.xlsx"
        reused = path.exists()
        if reused:
            raw = path.read_bytes()
            error = None
        else:
            raw, error = fetch_kap_excel(event_id)
            if raw is None:
                return {
                    "event_id": event_id,
                    "captured": False,
                    "error": error,
                }
            path.write_bytes(raw)

        if raw[:2] != b"PK":
            return {
                "event_id": event_id,
                "captured": False,
                "error": "EXCEL_CAPTURE_NOT_ZIP",
            }
        try:
            cells = parse_excel_cells(raw)
        except Exception as exc:  # noqa: BLE001
            return {
                "event_id": event_id,
                "captured": False,
                "error": f"EXCEL_PARSE:{type(exc).__name__}:{exc}",
            }

        parsed = {
            "event_id": event_id,
            "source_url": KAP_EXCEL_TEMPLATE.format(event_id=event_id),
            "raw_sha256": sha256_bytes(raw),
            "cells": cells,
        }
        parsed_path = parsed_dir / f"{event_id}.json"
        parsed_path.write_text(
            json.dumps(
                parsed,
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
            + "\n",
            encoding="utf-8",
        )
        return {
            "event_id": event_id,
            "captured": True,
            "error": None,
            "raw_sha256": parsed["raw_sha256"],
            "parsed_sha256": sha256_bytes(parsed_path.read_bytes()),
            "nonempty_cell_count": len(cells),
            "capture_source": (
                "REUSED_FROZEN_XLSX"
                if reused
                else "NEW_OFFICIAL_XLSX"
            ),
        }

    event_ids = queue["event_id"].astype(str).tolist()
    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
        future_map = {
            pool.submit(one, event_id): event_id
            for event_id in event_ids
        }
        for future in concurrent.futures.as_completed(future_map):
            rows.append(future.result())

    capture = pd.DataFrame(rows).sort_values("event_id").reset_index(drop=True)
    for row in capture.loc[capture["captured"].eq(True)].itertuples(index=False):
        parsed = json.loads(
            (parsed_dir / f"{row.event_id}.json").read_text(encoding="utf-8")
        )
        for cell in parsed["cells"]:
            cell_records.append(
                {
                    "event_id": str(row.event_id),
                    **cell,
                }
            )
    return capture, cell_records


def capture_details(
    queue: pd.DataFrame,
    *,
    fetch_missing: bool = False,
) -> tuple[pd.DataFrame, list[dict]]:
    detail_dir = OUT / "kap_detail"
    parsed_dir = OUT / "parsed_detail"
    detail_dir.mkdir(parents=True, exist_ok=True)
    parsed_dir.mkdir(parents=True, exist_ok=True)

    captured_at = datetime.now(UTC)
    if captured_at <= ACTIVATION:
        raise RuntimeError("45F official detail capture must be post activation")

    rows = []
    field_records: list[dict] = []

    def one(event_id: str) -> dict:
        raw_path = detail_dir / f"{event_id}.html.gz"
        reused = raw_path.exists()
        if reused:
            try:
                raw = gzip.decompress(raw_path.read_bytes())
                error = None
            except Exception as exc:
                return {
                    "event_id": event_id,
                    "captured": False,
                    "error": f"EXISTING_CAPTURE_CORRUPT:{type(exc).__name__}",
                }
        else:
            if not fetch_missing:
                return {
                    "event_id": event_id,
                    "captured": False,
                    "error": "HTML_NOT_CAPTURED_A2_EXCEL_PRIMARY",
                }
            raw, error = fetch_kap(event_id)
            if raw is None:
                return {
                    "event_id": event_id,
                    "captured": False,
                    "error": error,
                }

        expanded = decode_next_payload(raw)
        basic = parse_basic(expanded)
        fields = parse_taxonomy_fields(raw, expanded)
        if str(basic.get("disclosureIndex") or "") not in {"", event_id}:
            return {
                "event_id": event_id,
                "captured": False,
                "error": "DISCLOSURE_INDEX_CONFLICT",
            }

        if reused:
            raw_gzip_sha = sha256_bytes(raw_path.read_bytes())
        else:
            raw_gzip_sha = write_gzip(raw_path, raw)
        parsed = {
            "event_id": event_id,
            "captured_at": captured_at.isoformat(),
            "source_url": KAP_TEMPLATE.format(event_id=event_id),
            "raw_sha256": sha256_bytes(raw),
            "raw_gzip_sha256": raw_gzip_sha,
            "basic": basic,
            "fields": fields,
        }
        parsed_path = parsed_dir / f"{event_id}.json"
        parsed_path.write_text(
            json.dumps(parsed, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )

        return {
            "event_id": event_id,
            "captured": True,
            "error": None,
            "raw_sha256": parsed["raw_sha256"],
            "raw_gzip_sha256": raw_gzip_sha,
            "parsed_sha256": sha256_bytes(parsed_path.read_bytes()),
            "field_count": len(fields),
            "basic_title": basic.get("title"),
            "basic_publish_date": basic.get("publishDate"),
            "capture_source": (
                "REUSED_FROZEN_HTML" if reused else "NEW_OFFICIAL_HTML"
            ),
        }

    event_ids = queue["event_id"].astype(str).tolist()
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
        future_map = {pool.submit(one, event_id): event_id for event_id in event_ids}
        for future in concurrent.futures.as_completed(future_map):
            rows.append(future.result())

    capture = pd.DataFrame(rows).sort_values("event_id").reset_index(drop=True)
    for row in capture.loc[capture["captured"].eq(True)].itertuples(index=False):
        parsed = json.loads(
            (parsed_dir / f"{row.event_id}.json").read_text(encoding="utf-8")
        )
        for field in parsed["fields"]:
            field_records.append(
                {
                    "event_id": str(row.event_id),
                    "field_name": field["field_name"],
                    "titles_tr": "|".join(field["titles_tr"]),
                    "values_tr": "|".join(field["values_tr"]),
                }
            )

    return capture, field_records


def capture_yahoo_actions(tickers: list[str]) -> pd.DataFrame:
    rows: list[dict] = []
    end_exclusive = (datetime.now(UTC).date() + timedelta(days=1)).isoformat()

    for offset in range(0, len(tickers), 50):
        batch = tickers[offset : offset + 50]
        symbols = [f"{ticker}.IS" for ticker in batch]
        try:
            frame = yf.download(
                symbols,
                start=PRICE_START,
                end=end_exclusive,
                auto_adjust=False,
                actions=True,
                repair=False,
                progress=False,
                threads=True,
                group_by="ticker",
            )
        except Exception:
            continue
        if frame.empty or not isinstance(frame.columns, pd.MultiIndex):
            continue

        available = set(frame.columns.get_level_values(0))
        for ticker in batch:
            symbol = f"{ticker}.IS"
            if symbol not in available:
                continue
            group = frame[symbol]
            for index, row in group.iterrows():
                dividend = float(row.get("Dividends") or 0.0)
                split = float(row.get("Stock Splits") or 0.0)
                if dividend == 0.0 and split == 0.0:
                    continue
                rows.append(
                    {
                        "ticker": ticker,
                        "source_symbol": symbol,
                        "trade_date": pd.Timestamp(index).date().isoformat(),
                        "dividend": dividend,
                        "stock_split": split,
                    }
                )

    if not rows:
        return pd.DataFrame(
            columns=[
                "ticker",
                "source_symbol",
                "trade_date",
                "dividend",
                "stock_split",
            ]
        )
    return (
        pd.DataFrame(rows)
        .sort_values(["ticker", "trade_date"])
        .drop_duplicates(["ticker", "trade_date"])
        .reset_index(drop=True)
    )

def vendor_corroboration(
    queue: pd.DataFrame,
    actions: pd.DataFrame,
) -> dict[str, bool]:
    by_ticker: dict[str, pd.DataFrame] = {
        ticker: group.copy()
        for ticker, group in actions.groupby("ticker", sort=False)
    }
    out: dict[str, bool] = {}
    for row in queue.itertuples(index=False):
        try:
            published = pd.to_datetime(row.publish_date, dayfirst=True).date()
        except Exception:
            out[str(row.event_id)] = False
            continue
        matched = False
        for ticker in str(row.tickers).split("|"):
            group = by_ticker.get(ticker)
            if group is None:
                continue
            dates = pd.to_datetime(group["trade_date"]).dt.date
            distance = dates.map(lambda day: abs((day - published).days))
            if bool((distance <= 45).any()):
                matched = True
                break
        out[str(row.event_id)] = matched
    return out


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    queue = load_queue()
    universe = pd.read_csv(UNIVERSE, dtype=str)
    tickers = sorted(set(universe["ticker"].astype(str)))
    if len(tickers) != 100:
        raise RuntimeError("45F expected exact 100-name universe")

    prior_provenance_path = OUT / "provenance.json"
    prior_provenance = (
        json.loads(prior_provenance_path.read_text(encoding="utf-8"))
        if prior_provenance_path.exists()
        else None
    )

    capture, field_records = capture_details(queue, fetch_missing=False)
    excel_capture, excel_cells = capture_excel_details(queue)

    actions_path = OUT / "yahoo_actions.csv"
    if actions_path.exists():
        actions = pd.read_csv(actions_path)
    else:
        actions = capture_yahoo_actions(tickers)
    vendor = vendor_corroboration(queue, actions)

    queue_path = OUT / "event_queue.csv"
    queue.to_csv(queue_path, index=False, lineterminator="\n")

    capture_path = OUT / "detail_capture.csv"
    capture.to_csv(capture_path, index=False, lineterminator="\n")

    fields_path = OUT / "taxonomy_fields.csv"
    pd.DataFrame(field_records).to_csv(
        fields_path,
        index=False,
        lineterminator="\n",
    )

    excel_capture_path = OUT / "excel_capture.csv"
    excel_capture.to_csv(
        excel_capture_path,
        index=False,
        lineterminator="\n",
    )

    excel_cells_path = OUT / "excel_cells.csv.gz"
    excel_cells_frame = pd.DataFrame(excel_cells)
    excel_cells_frame.to_csv(
        excel_cells_path,
        index=False,
        lineterminator="\n",
        compression={
            "method": "gzip",
            "compresslevel": 9,
            "mtime": 0,
        },
    )

    actions_path = OUT / "yahoo_actions.csv"
    actions.to_csv(actions_path, index=False, lineterminator="\n")

    initial_rows = []
    capture_by_id = {
        str(row.event_id): row
        for row in capture.itertuples(index=False)
    }
    excel_by_id = {
        str(row.event_id): row
        for row in excel_capture.itertuples(index=False)
    }
    for row in queue.itertuples(index=False):
        event_id = str(row.event_id)
        cap = capture_by_id[event_id]
        xls = excel_by_id[event_id]
        html_captured = bool(cap.captured)
        excel_captured = bool(xls.captured)
        captured = html_captured or excel_captured
        official_sha = (
            str(cap.raw_sha256)
            if html_captured
            else (str(xls.raw_sha256) if excel_captured else None)
        )
        evidence = EconomicResolutionEvidence(
            event_id=event_id,
            event_type=str(row.event_type),
            official_detail_captured=captured,
            official_detail_sha256=official_sha,
            # First 45F run is deliberately field-discovery only. Economic
            # completeness is promoted only after explicit field contracts
            # are written from captured evidence.
            official_economic_fields_complete=False,
            official_non_price_affecting_explicit=False,
            vendor_corroboration_present=bool(vendor.get(event_id, False)),
            source_conflict=False,
        )
        result = resolve_economic_action(evidence)
        initial_rows.append(
            {
                "event_id": event_id,
                "event_type": row.event_type,
                "tickers": row.tickers,
                "scopes": row.scopes,
                "official_detail_captured": captured,
                "official_detail_sha256": evidence.official_detail_sha256,
                "official_html_captured": html_captured,
                "official_excel_captured": excel_captured,
                "html_sha256": (
                    str(cap.raw_sha256) if html_captured else None
                ),
                "excel_sha256": (
                    str(xls.raw_sha256) if excel_captured else None
                ),
                "taxonomy_field_count": (
                    int(cap.field_count) if html_captured else 0
                ),
                "excel_nonempty_cell_count": (
                    int(xls.nonempty_cell_count)
                    if excel_captured
                    else 0
                ),
                "vendor_corroboration_present": evidence.vendor_corroboration_present,
                "status": result.status.value,
                "risk_released": result.risk_released,
                "reason_codes": "|".join(result.reason_codes),
                "shadow_signal_allowed": result.shadow_signal_allowed,
            }
        )

    initial = pd.DataFrame(initial_rows).sort_values("event_id").reset_index(drop=True)
    if initial["risk_released"].astype(bool).any():
        raise RuntimeError(
            "field-discovery run cannot release event risk before field contracts"
        )
    if initial["shadow_signal_allowed"].astype(bool).any():
        raise RuntimeError("45F cannot authorize shadow signals")

    resolution_path = OUT / "initial_resolution.csv"
    initial.to_csv(resolution_path, index=False, lineterminator="\n")

    fields_frame = pd.DataFrame(field_records)
    if fields_frame.empty:
        field_summary = []
    else:
        event_type_by_id = dict(zip(queue["event_id"].astype(str), queue["event_type"], strict=True))
        fields_frame["event_type"] = fields_frame["event_id"].map(event_type_by_id)
        summary = (
            fields_frame.groupby(["event_type", "field_name"], sort=True)
            .agg(
                event_count=("event_id", "nunique"),
                nonempty_value_count=("values_tr", lambda s: int((s.astype(str).str.len() > 0).sum())),
            )
            .reset_index()
        )
        field_summary = summary.to_dict("records")

    now = datetime.now(UTC).isoformat()
    if excel_cells_frame.empty:
        excel_inventory = []
    else:
        event_type_by_id = dict(
            zip(
                queue["event_id"].astype(str),
                queue["event_type"],
                strict=True,
            )
        )
        excel_cells_frame["event_type"] = excel_cells_frame["event_id"].map(
            event_type_by_id
        )
        excel_summary = (
            excel_cells_frame.groupby("event_type", sort=True)
            .agg(
                event_count=("event_id", "nunique"),
                nonempty_cell_count=("value", "count"),
            )
            .reset_index()
        )
        excel_inventory = excel_summary.to_dict("records")

    provenance = {
        "contract": "CORPORATE_ACTION_ECONOMIC_RESOLUTION_FIELD_DISCOVERY_V1",
        "captured_at": (
            prior_provenance.get("captured_at")
            if prior_provenance
            else now
        ),
        "last_repair_at": now,
        "resumable_capture": True,
        "production_ready": False,
        "score_values_computed": False,
        "real_shadow_run_created": False,
        "queue": {
            "ticker_event_rows": 411,
            "unique_events": int(len(queue)),
            "affected_tickers": int(
                len(
                    set(
                        ticker
                        for value in queue["tickers"].astype(str)
                        for ticker in value.split("|")
                        if ticker
                    )
                )
            ),
        },
        "official_detail": {
            "captured": int(capture["captured"].astype(bool).sum()),
            "failed": int((~capture["captured"].astype(bool)).sum()),
            "field_records": int(len(field_records)),
            "reused_frozen_html": int(
                capture.get("capture_source", pd.Series(dtype=str))
                .astype(str)
                .eq("REUSED_FROZEN_HTML")
                .sum()
            ),
            "new_official_html": int(
                capture.get("capture_source", pd.Series(dtype=str))
                .astype(str)
                .eq("NEW_OFFICIAL_HTML")
                .sum()
            ),
        },
        "official_excel": {
            "captured": int(
                excel_capture["captured"].astype(bool).sum()
            ),
            "failed": int(
                (~excel_capture["captured"].astype(bool)).sum()
            ),
            "cell_records": int(len(excel_cells)),
            "reused_frozen_xlsx": int(
                excel_capture.get(
                    "capture_source",
                    pd.Series(dtype=str),
                )
                .astype(str)
                .eq("REUSED_FROZEN_XLSX")
                .sum()
            ),
            "new_official_xlsx": int(
                excel_capture.get(
                    "capture_source",
                    pd.Series(dtype=str),
                )
                .astype(str)
                .eq("NEW_OFFICIAL_XLSX")
                .sum()
            ),
            "event_type_inventory": excel_inventory,
        },
        "vendor_corroboration": {
            "yahoo_action_rows": int(len(actions)),
            "events_with_nearby_vendor_action": int(sum(vendor.values())),
            "can_resolve_event_alone": False,
        },
        "initial_resolution": {
            str(key): int(value)
            for key, value in initial["status"].value_counts().sort_index().items()
        },
        "risk_released": 0,
        "field_inventory": field_summary,
        "next_required_work": (
            "LOCK_ACTION_TYPE_SPECIFIC_OFFICIAL_FIELD_CONTRACTS_FROM_CAPTURED_TAXONOMY"
        ),
    }
    provenance_path = OUT / "provenance.json"
    provenance_path.write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )

    files = [
        queue_path,
        capture_path,
        fields_path,
        excel_capture_path,
        excel_cells_path,
        actions_path,
        resolution_path,
        provenance_path,
    ]
    files.extend(path for path in (OUT / "kap_detail").glob("*") if path.is_file())
    files.extend(
        path for path in (OUT / "parsed_detail").glob("*") if path.is_file()
    )
    files.extend(
        path for path in (OUT / "kap_excel").glob("*") if path.is_file()
    )
    files.extend(
        path for path in (OUT / "parsed_excel").glob("*") if path.is_file()
    )
    sums = [
        f"{sha256_bytes(path.read_bytes())}  {path.relative_to(OUT)}"
        for path in sorted(files, key=lambda item: str(item.relative_to(OUT)))
    ]
    (OUT / "SHA256SUMS").write_text("\n".join(sums) + "\n", encoding="utf-8")

    print(json.dumps(provenance, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
